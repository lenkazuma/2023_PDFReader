import sys
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document as DocxDocument
from langchain_core.documents import Document
from langchain_core.embeddings import DeterministicFakeEmbedding
from langchain_core.language_models.fake_chat_models import FakeListChatModel
from pypdf import PdfWriter

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from reader_core import (  # noqa: E402
    DOCX_MIME,
    PDF_MIME,
    EmptyDocumentError,
    answer_question,
    build_index,
    extract_documents,
    file_fingerprint,
    split_documents,
    summarize,
)


def make_docx(paragraphs, table_rows=()):
    doc = DocxDocument()
    for text in paragraphs:
        doc.add_paragraph(text)
    if table_rows:
        table = doc.add_table(rows=len(table_rows), cols=len(table_rows[0]))
        for r, row in enumerate(table_rows):
            for c, value in enumerate(row):
                table.cell(r, c).text = value
    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def test_docx_extraction_includes_tables():
    data = make_docx(["Revenue grew strongly."], [("Year", "Revenue"), ("2024", "120")])
    docs = extract_documents(data, DOCX_MIME)
    assert "Revenue grew strongly." in docs[0].page_content
    assert "2024 | 120" in docs[0].page_content


def test_blank_pdf_raises_empty_document_error():
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    buffer = BytesIO()
    writer.write(buffer)
    with pytest.raises(EmptyDocumentError):
        extract_documents(buffer.getvalue(), PDF_MIME)


def test_unsupported_format_is_rejected():
    with pytest.raises(ValueError):
        extract_documents(b"hello", "text/plain", "notes.txt")


def test_split_keeps_page_metadata():
    docs = [Document(page_content="word " * 400, metadata={"page": 7})]
    chunks = split_documents(docs, chunk_size=300, chunk_overlap=50)
    assert len(chunks) > 1
    assert all(chunk.metadata["page"] == 7 for chunk in chunks)


def test_fingerprint_is_stable_and_content_based():
    assert file_fingerprint(b"a") == file_fingerprint(b"a")
    assert file_fingerprint(b"a") != file_fingerprint(b"b")


def test_answer_question_returns_sources_and_includes_history():
    chunks = [
        Document(page_content="The revenue was 120 million.", metadata={"page": 2}),
        Document(page_content="The CEO is Alice.", metadata={"page": 5}),
    ]
    index = build_index(chunks, DeterministicFakeEmbedding(size=16))
    llm = FakeListChatModel(responses=["120 million (p. 2)"])
    result = answer_question("What was revenue?", index, llm, history=[("Hi", "Hello")], k=2)
    assert result.text == "120 million (p. 2)"
    assert {doc.metadata["page"] for doc in result.sources} == {2, 5}


def test_summarize_single_call_for_small_documents():
    llm = FakeListChatModel(responses=["short summary"])
    chunks = [Document(page_content="tiny", metadata={"page": 1})]
    assert summarize(chunks, llm) == "short summary"


def test_summarize_map_reduce_for_large_documents():
    llm = FakeListChatModel(responses=["part A", "part B", "part C", "final"])
    chunks = [Document(page_content="x" * 50, metadata={"page": i}) for i in range(3)]
    assert summarize(chunks, llm, max_chars_per_call=60) == "final"
