"""Document parsing, indexing and question answering, independent of Streamlit."""

from __future__ import annotations

import hashlib
from dataclasses import dataclass
from io import BytesIO

from docx import Document as DocxDocument
from langchain_core.documents import Document
from langchain_core.embeddings import Embeddings
from langchain_core.language_models import BaseChatModel
from langchain_core.vectorstores import InMemoryVectorStore
from langchain_text_splitters import RecursiveCharacterTextSplitter
from pypdf import PdfReader

PDF_MIME = "application/pdf"
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

QA_SYSTEM_PROMPT = (
    "You answer questions about a document using only the excerpts provided. "
    "Each excerpt is labelled with its page. Cite pages like (p. 3). "
    "If the excerpts do not contain the answer, say you could not find it. "
    "Reply in the same language as the question."
)
SUMMARY_PROMPT = (
    "Write a concise summary of the following document content in the language the document is written in. "
    "Focus on the main points.\n\n{text}"
)
COMBINE_PROMPT = (
    "Combine these partial summaries of one document into a single concise summary, "
    "in the language they are written in.\n\n{text}"
)


class EmptyDocumentError(ValueError):
    """Raised when no text could be extracted (e.g. a scanned PDF)."""


@dataclass
class Answer:
    text: str
    sources: list[Document]


def file_fingerprint(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def extract_documents(data: bytes, mime_type: str, filename: str = "") -> list[Document]:
    """Return one Document per PDF page, or one for the whole DOCX (paragraphs + tables)."""
    if mime_type == PDF_MIME or filename.lower().endswith(".pdf"):
        reader = PdfReader(BytesIO(data))
        docs = [
            Document(page_content=text, metadata={"page": number})
            for number, page in enumerate(reader.pages, start=1)
            if (text := (page.extract_text() or "").strip())
        ]
    elif mime_type == DOCX_MIME or filename.lower().endswith(".docx"):
        doc = DocxDocument(BytesIO(data))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        text = "\n".join(parts).strip()
        docs = [Document(page_content=text, metadata={"page": 1})] if text else []
    else:
        raise ValueError("Unsupported file format. Please upload a PDF or DOCX file.")

    if not docs:
        raise EmptyDocumentError(
            "No text could be extracted. The file may be a scanned image; try a PDF with selectable text."
        )
    return docs


def split_documents(docs: list[Document], chunk_size: int = 1000, chunk_overlap: int = 200) -> list[Document]:
    splitter = RecursiveCharacterTextSplitter(chunk_size=chunk_size, chunk_overlap=chunk_overlap)
    return splitter.split_documents(docs)


def build_index(chunks: list[Document], embeddings: Embeddings) -> InMemoryVectorStore:
    return InMemoryVectorStore.from_documents(chunks, embeddings)


def _format_excerpts(docs: list[Document]) -> str:
    return "\n\n".join(f"[p. {doc.metadata.get('page', '?')}]\n{doc.page_content}" for doc in docs)


def answer_question(
    question: str,
    index: InMemoryVectorStore,
    llm: BaseChatModel,
    history: list[tuple[str, str]] | None = None,
    k: int = 4,
) -> Answer:
    sources = index.similarity_search(question, k=k)
    messages: list[tuple[str, str]] = [("system", QA_SYSTEM_PROMPT)]
    for past_question, past_answer in (history or [])[-4:]:
        messages += [("human", past_question), ("ai", past_answer)]
    messages.append(("human", f"Excerpts:\n{_format_excerpts(sources)}\n\nQuestion: {question}"))
    return Answer(text=str(llm.invoke(messages).content), sources=sources)


def summarize(chunks: list[Document], llm: BaseChatModel, max_chars_per_call: int = 12000) -> str:
    """Stuff small documents into one call; map-reduce larger ones."""
    groups: list[str] = []
    current = ""
    for chunk in chunks:
        if current and len(current) + len(chunk.page_content) > max_chars_per_call:
            groups.append(current)
            current = ""
        current += chunk.page_content + "\n\n"
    if current:
        groups.append(current)

    if len(groups) == 1:
        return str(llm.invoke(SUMMARY_PROMPT.format(text=groups[0])).content)

    partials = [str(result.content) for result in llm.batch([SUMMARY_PROMPT.format(text=g) for g in groups])]
    while len(partials) > 1:
        combined = "\n\n".join(partials)
        if len(combined) <= max_chars_per_call:
            return str(llm.invoke(COMBINE_PROMPT.format(text=combined)).content)
        half = len(partials) // 2
        partials = [
            str(llm.invoke(COMBINE_PROMPT.format(text="\n\n".join(part))).content)
            for part in (partials[:half], partials[half:])
        ]
    return partials[0]
